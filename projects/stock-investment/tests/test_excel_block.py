"""_write_block が書き出すレイアウトを固定するテスト。

1ブロック（実行時刻ヘッダー → 列ヘッダー → ランキング → 総評）の
行の進み方とセルの中身を、現行の出力そのままに写している。
切り出しリファクタで行番号がずれると後続ブロックが上書きされるため、
戻り値（次に書き始める行）も併せて固定する。

ファイルには書き出さない。openpyxl のワークブックをメモリ上で組み立てる。
"""

from datetime import datetime

import openpyxl
import pytest

from src.report.excel_exporter import _write_block

NOW = datetime(2026, 9, 1, 14, 30)

RANKINGS = [
    {"rank": 1, "ticker": "7203", "name": "トヨタ", "stars": "★★★★★",
     "confidence": 85, "close": 3100.0, "RSI14": 52.3, "MACD方向": "上昇",
     "SMA20比": "-2.1%", "BB位置": "中間", "STOCH_K": 45.0,
     "reason": "押し目", "news_basis": "決算好調"},
    {"rank": 2, "ticker": "6758", "name": "ソニー", "stars": "★★★",
     "confidence": 70, "close": 14000.0, "RSI14": 61.0, "MACD方向": "横ばい",
     "SMA20比": "+1.0%", "BB位置": "上限", "STOCH_K": 80.0,
     "reason": "様子見", "news_basis": "特になし"},
]


def _result(**over):
    base = {
        "flow": "通常",
        "market": "東証プライム",
        "vix": 18.2,
        "fear_greed": 55,
        "rankings": RANKINGS,
        "summary": "全体に落ち着いた地合い。",
    }
    base.update(over)
    return base


@pytest.fixture
def ws():
    return openpyxl.Workbook().active


def test_block_row_layout_and_return_value(ws):
    """1ブロックは 実行時刻+列ヘッダー+ランキング数+総評2行 を消費する。"""
    next_row = _write_block(ws, _result(), start_row=1, now=NOW)
    assert next_row == 1 + 1 + 1 + len(RANKINGS) + 2 == 7


def test_block_can_start_partway_down(ws):
    """同日2回目の追記に備えて、開始行をずらしても同じだけ進む。"""
    assert _write_block(ws, _result(), start_row=10, now=NOW) == 16


def test_timestamp_header_text(ws):
    _write_block(ws, _result(), start_row=1, now=NOW)
    assert ws.cell(row=1, column=1).value == (
        "実行: 14:30　【通常】　東証プライム  VIX:18.2  恐怖&欲指数:55/100"
    )
    assert ws.row_dimensions[1].height == 22


def test_sentiment_is_omitted_when_absent(ws):
    _write_block(ws, _result(vix=None, fear_greed=None), start_row=1, now=NOW)
    assert ws.cell(row=1, column=1).value == "実行: 14:30　【通常】　東証プライム"


def test_flow_label_is_translated(ws):
    """FLOW_LABELS にある値は表示名に置き換える。"""
    _write_block(ws, _result(flow="ニュース起点"), start_row=1, now=NOW)
    assert "【ニュース分析】" in ws.cell(row=1, column=1).value


def test_column_headers_without_news_column(ws):
    _write_block(ws, _result(), start_row=1, now=NOW)
    headers = [ws.cell(row=2, column=c).value for c in range(1, 14)]
    assert headers == [
        "順位", "コード", "銘柄名", "おすすめ度", "AIの\n確信度(%)", "株価\n(終値)",
        "売買タイミング\n(RSI)", "トレンド方向\n(MACD)", "20日平均\n株価比",
        "価格帯の\n位置(BB)", "短期過熱度\n(Stoch)", "おすすめ理由と売買タイミング",
        None,
    ]
    assert ws.row_dimensions[2].height == 32


def test_news_column_appears_only_for_news_flows(ws):
    _write_block(ws, _result(flow="YouTube起点"), start_row=1, now=NOW)
    assert ws.cell(row=2, column=13).value == "参考にしたニュース・動画"
    assert ws.cell(row=3, column=13).value == "決算好調"


def test_ranking_rows_contents(ws):
    _write_block(ws, _result(), start_row=1, now=NOW)
    assert [ws.cell(row=3, column=c).value for c in range(1, 13)] == [
        1, "7203", "トヨタ", "★★★★★", 85, 3100.0, 52.3, "上昇", "-2.1%",
        "中間", 45.0, "押し目",
    ]
    assert [ws.cell(row=4, column=c).value for c in range(1, 13)] == [
        2, "6758", "ソニー", "★★★", 70, 14000.0, 61.0, "横ばい", "+1.0%",
        "上限", 80.0, "様子見",
    ]
    assert ws.row_dimensions[3].height == 45


def test_summary_section(ws):
    _write_block(ws, _result(), start_row=1, now=NOW)
    assert ws.cell(row=5, column=1).value == "【まとめ・市場の状況】"
    assert ws.cell(row=6, column=1).value == "全体に落ち着いた地合い。"
    assert ws.row_dimensions[6].height == 65


def test_market_outlook_is_appended_to_summary(ws):
    _write_block(ws, _result(market_outlook="円安が続く"), start_row=1, now=NOW)
    assert ws.cell(row=6, column=1).value == (
        "全体に落ち着いた地合い。\n【市場全体の状況】円安が続く"
    )


def test_empty_rankings_still_writes_summary(ws):
    next_row = _write_block(ws, _result(rankings=[]), start_row=1, now=NOW)
    assert next_row == 5
    assert ws.cell(row=3, column=1).value == "【まとめ・市場の状況】"

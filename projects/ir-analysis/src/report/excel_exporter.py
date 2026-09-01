from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPORT_DIR = Path(__file__).parent.parent.parent / "data" / "reports"

IMPACT_COLORS = {
    "positive": "C6EFCE",
    "negative": "FFC7CE",
    "neutral":  "FFEB9C",
}

RELEVANCE_COLORS = {
    "high":   "FF0000",
    "medium": "FF8C00",
    "low":    "888888",
}


def _header_style(ws, row: int, headers: list[str]):
    fill = PatternFill("solid", fgColor="2F4F8F")
    font = Font(bold=True, color="FFFFFF", size=11)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = fill
        c.font = font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export_to_excel(items: list[dict], output_path: str | None = None) -> str:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    if output_path is None:
        today = date.today().strftime("%Y%m%d")
        output_path = str(REPORT_DIR / f"ir_report_{today}.xlsx")

    wb = openpyxl.Workbook()

    # ── シート1: 分析サマリー ──
    ws = wb.active
    ws.title = "IR分析サマリー"

    headers = [
        "企業名", "銘柄コード", "開示タイトル", "開示日",
        "要約", "影響", "影響理由",
        "スイング重要度", "スイングコメント",
        "売上", "営業利益", "修正内容",
        "重要ポイント",
    ]
    _header_style(ws, 1, headers)
    ws.row_dimensions[1].height = 30

    col_widths = [16, 8, 35, 10, 40, 10, 35, 12, 35, 12, 12, 15, 50]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w

    for row_idx, item in enumerate(items, 2):
        a = item.get("analysis", {})
        nums = a.get("numbers", {})
        impact = a.get("impact", "neutral")
        relevance = a.get("swing_relevance", "low")

        values = [
            item.get("company", ""),
            item.get("code", ""),
            item.get("title", ""),
            item.get("date", ""),
            a.get("summary", ""),
            impact,
            a.get("impact_reason", ""),
            relevance,
            a.get("swing_note", ""),
            nums.get("売上", ""),
            nums.get("営業利益", ""),
            nums.get("修正内容", ""),
            " / ".join(a.get("key_points", [])),
        ]

        for col, val in enumerate(values, 1):
            c = ws.cell(row=row_idx, column=col, value=val)
            c.alignment = Alignment(wrap_text=True, vertical="top")

        # 影響度の背景色
        impact_color = IMPACT_COLORS.get(impact)
        if impact_color:
            fill = PatternFill("solid", fgColor=impact_color)
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = fill

        # スイング重要度の文字色
        rel_color = RELEVANCE_COLORS.get(relevance)
        if rel_color:
            ws.cell(row=row_idx, column=8).font = Font(bold=True, color=rel_color)

        ws.row_dimensions[row_idx].height = 60

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ── シート2: 凡例 ──
    ws2 = wb.create_sheet("凡例")
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 40

    legend = [
        ("【影響度】", ""),
        ("positive（緑）", "株価にプラスの材料"),
        ("negative（赤）", "株価にマイナスの材料"),
        ("neutral（黄）",  "中立・影響軽微"),
        ("", ""),
        ("【スイング重要度】", ""),
        ("high（赤）",   "スイングで注目すべき開示"),
        ("medium（橙）", "参考程度"),
        ("low（灰）",    "スイングへの影響軽微"),
    ]
    for r, (a, b) in enumerate(legend, 1):
        ws2.cell(r, 1, a).font = Font(bold=True) if "【" in a else Font()
        ws2.cell(r, 2, b)

    wb.save(output_path)
    return output_path
